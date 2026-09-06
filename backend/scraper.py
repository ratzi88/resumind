import sys
import os
from pypdf import PdfReader
from docx import Document


def get_file_extension(filepath):
    return os.path.splitext(filepath)[1].lower()


def scrape_file(file_path, extension, ocr_lang="eng"):
    """
    Extract plain text from a PDF or DOCX file.
    Falls back to OCR (pytesseract) for scanned PDF pages.
    """
    if extension == ".pdf":
        text_parts = []
        ocr_needed_pages = []

        reader = PdfReader(file_path)
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(page_text)
            else:
                text_parts.append(None)
                ocr_needed_pages.append(i)

        if ocr_needed_pages:
            try:
                from pdf2image import convert_from_path
                import pytesseract
            except ImportError as exc:
                raise RuntimeError(
                    "This PDF appears to be scanned. Install pdf2image and "
                    "pytesseract, and make sure the Tesseract binary is available."
                ) from exc

            print(f"OCR needed on {len(ocr_needed_pages)} page(s)...")
            images = convert_from_path(file_path, dpi=300)
            for i in ocr_needed_pages:
                ocr_text = pytesseract.image_to_string(images[i], lang=ocr_lang)
                text_parts[i] = ocr_text or ""

        text = "\n".join(tp for tp in text_parts if tp is not None)
        return text.strip()

    elif extension == ".docx":
        doc = Document(file_path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    text += "\n" + " | ".join(cells)
        return text.strip()

    else:
        raise ValueError("Unsupported file format. Please provide a PDF or DOCX file.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python scraper.py <file_path>")
        sys.exit(1)

    file_path = sys.argv[1]
    extension = get_file_extension(file_path)

    print(f"Processing file: {file_path} ({extension})\n")
    scraped_text = scrape_file(file_path, extension)
    print("=== Extracted Text ===\n")
    print(scraped_text)
